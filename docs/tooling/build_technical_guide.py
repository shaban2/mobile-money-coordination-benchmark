"""Build the onboarding DOCX from TECHNICAL_GUIDE.md, without touching evidence.

Requires python-docx and Pillow. Use the Codex bundled Python when available.
The Markdown contains Mermaid for repository readers; the DOCX uses compact
equivalent flow diagrams. Source excerpts and the symbol index are materialized
in Markdown before building, so the two deliverables contain the same text.
"""
from pathlib import Path
import re
import tempfile
import textwrap

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/TECHNICAL_GUIDE.md'
OUTPUT = ROOT / 'docs/mobile-money-coordination-experiment-technical-documentation.docx'

# File descriptions are authored; symbol locations are checked against source.
FILES = [
 ('src/server.js', 'HTTP entry point; reads environment, starts system, routes API/admin requests, shuts down on signals.'),
 ('src/system.js', 'createSystem wires persistence, bus, coordinator, adapters and application; start/recover/reset/stop control their lifecycle.'),
 ('src/network-gateway.js', 'selectUpstream routes API or callback traffic; streaming HTTP proxy emits request IDs and timing/error logs.'),
 ('src/adapter-server.js', 'Standalone HTTP entry point constructing the profile router, local boundary and journal-backed simulator.'),
 ('src/domain/transfer.js', 'validateTransferCommand rejects invalid fields; normalizeTransferCommand and commandFingerprint define replay identity; createTransferRecord assigns IDs; transitionTransfer enforces states; publicTransfer shapes responses.'),
 ('src/domain/errors.js', 'DomainError carries public error code and HTTP status; asErrorBody hides unexpected internal failures.'),
 ('src/application/payment-application.js', 'submit accepts/replays; process tracks background work; recover reloads candidates; get returns status; drain waits for tasks; invariants reconciles ledger and state.'),
 ('src/application/callback-sink.js', 'CallbackSink.deliver sends the terminal public record, retries configured attempts and persists delivery evidence.'),
 ('src/coordination/rest-orchestrator.js', 'RestOrchestrator.execute advances direct workflow steps within WorkLimiter; calls shared provider helper and persistence.'),
 ('src/coordination/kafka-choreography.js', 'KafkaChoreography.registerHandlers connects topic stages; execute initiates/resumes work and waits for terminal notification; transition coordinates inbox/outbox state updates.'),
 ('src/coordination/workflow-steps.js', 'callProvider coordinates cached command results, common boundary, retry/deadline handling, provider attempts and uncertain outcomes; saveTransition and failTransfer are exported helper functions.'),
 ('src/adapters/provider-adapter.js', 'ProviderAdapter defines mapping methods; ProviderAdapterRouter.forProfile selects the registered A/B adapter.'),
 ('src/adapters/profile-a-adapter.js', 'ProfileAAdapter.toProviderRequest maps canonical fields to A; fromProviderResponse normalizes status and provider/failure references.'),
 ('src/adapters/profile-b-adapter.js', 'ProfileBAdapter supplies the different B field names and nested amount; response normalization preserves common workflow semantics.'),
 ('src/adapters/in-process-adapter-boundary.js', 'execute calls router, request mapper, simulator and response mapper; setAvailable supports local test faults.'),
 ('src/adapters/http-adapter-boundary.js', 'execute POSTs the transfer record to the common service; transport uncertainty gets an explicit error code; local availability control is rejected for this remote boundary.'),
 ('src/providers/provider-simulator.js', 'ProviderSimulator.submit simulates delay/outcome and idempotent effects; constructor restores the optional durable journal after restart.'),
 ('src/infrastructure/in-memory-persistence.js', 'Implements the application persistence interface using repository/ledger collections for fast tests; not durable PostgreSQL or a broker-backed outbox.'),
 ('src/infrastructure/in-memory-transfer-repository.js', 'InMemoryTransferRepository stores records and idempotency mappings; createOrGet/save/get/list support memory application flows.'),
 ('src/infrastructure/in-memory-ledger.js', 'InMemoryLedger stores balances and synthetic entries; used by memory persistence and reconciliation tests.'),
 ('src/infrastructure/in-memory-event-bus.js', 'InMemoryEventBus registers handlers and dispatches published events locally; a Kafka-coordinator unit test can use this substitute.'),
 ('src/infrastructure/kafka-event-bus.js', 'KafkaEventBus starts namespaced topics, producer and group consumer; subscribe/publish connect handlers; persistence inbox checks protect redelivery.'),
 ('src/infrastructure/outbox-dispatcher.js', 'start schedules polling; dispatchOnce prevents overlap; dispatchBatch publishes due rows and records success/backoff; stop waits for dispatch.'),
 ('src/infrastructure/terminal-notifier.js', 'TerminalNotifier.wait and notify connect durable terminal state to waiting workflow promises, with timeout cleanup.'),
 ('src/infrastructure/work-limiter.js', 'WorkLimiter.run bounds concurrently active tasks and releases queued work when a task settles.'),
 ('src/infrastructure/metrics.js', 'Metrics.increment/observe collect runtime counters and values; snapshot returns JSON; toPrometheus formats counters, sums/counts and gauges; reset clears them.'),
 ('src/infrastructure/runtime-diagnostics.js', 'RuntimeDiagnostics measures monotonic spans and runtime samples in bounded buffers; instrumentMethods wraps selected methods without changing their return/error semantics.'),
 ('src/infrastructure/postgres/postgres-persistence.js', 'PostgresPersistence is the SQL transaction boundary. Sections 6 and 10 trace createOrGet, transition, fulfill, confirmTerminal, provider/callback recording, outbox, recoveryCandidates and reset.'),
 ('src/infrastructure/postgres/migrate.js', 'migrate reads ordered SQL, skips recorded versions and applies each new migration atomically; called by persistence.start and the migration CLI.'),
 ('src/infrastructure/postgres/diagnostic-pool.js', 'instrumentPool wraps pool acquisition/query/release for diagnostic timing while preserving the pg API forms.'),
 ('src/infrastructure/postgres/database-diagnostics.js', 'DatabaseDiagnostics samples through a separate read-only observer, including wait/IO/WAL/checkpoint counters without exporting SQL text.'),
 ('src/experiment/pilot.js', 'evaluatePilotRun chooses protocol assessment; nextPilotStep selects bounded next work or review; makePilotBatch constructs paired rows with deterministic order.'),
 ('src/experiment/pilot-session.js', 'newPilotState, atPairBoundary, assertPilotCanStart/Budget/Evidence and executePilotSession implement durable sessions, pause/resume and cleanup/error disposition.'),
 ('src/experiment/pilot-timing.js', 'pilotTimingForPhase chooses short screening or full timing; pilotLifecycleMatches validates manifests; pilotTimingEnvironment formats runner environment.'),
 ('src/experiment/calibration.js', 'validateCalibrationProtocol and evaluateCalibrationRun enforce the fixed diagnostic calibration, not capacity inference.'),
 ('src/experiment/capacity-queue.js', 'assertCapacityRuleLaunchReady holds the draft classifier; exact observation batch/protocol checks authorize only reviewed descriptors; evaluateQueueCapacityRun validates evidence and correctness.'),
 ('src/experiment/queue-observation.js', 'Loads approved protocols; exploratoryProtocol and observationPairScope select a fixed scenario; describeQueueObservation reports finite sampled patterns without capacity pass/fail.'),
 ('src/experiment/exploratory-review.js', 'evidenceInventory hashes archived evidence; inspectCompletedObservationPair verifies it; predecessor/root review guards chain approved pairs.'),
 ('src/experiment/compose-lifecycle.js', 'resourceServicesFor defines expected services; composeArguments scopes Docker commands; stopComposeProject preserves volumes and verifies teardown; assertRunIsolation checks inventory.'),
 ('src/experiment/pilot-isolation-guard.js', 'observePilotIsolation listens to Docker starts; foreignContainerStart detects unrelated activity and invokes abort/report hooks.'),
 ('src/experiment/pilot-containers.js', 'assertRestorablePilotContainer checks exact approved identity and safe restoration, including rejecting auto-remove targets.'),
 ('src/experiment/provenance.js', 'sourceFiles defines the runtime inventory; fileHash and sourceSnapshotSha256 fingerprint bytes and relative paths for freezes and end checks.'),
 ('src/experiment/controller-diagnostics.js', 'captureCommandToFile preserves large stdout/stderr; recordControllerFailure retains the primary error; collectFailureDiagnostics records secondary collection problems.'),
 ('src/experiment/deep-diagnostic-evidence.js', 'assertDeepDiagnosticEvidence rejects missing/invalid deep database or Linux observations for protocols requiring them.'),
 ('src/experiment/host-resource-sampler.js', 'createHostResourceSampler captures host scope/counters; used by collect-docker-stats, not the transfer workflow.'),
 ('src/experiment/linux-resource-sampler.js', 'linuxDiagnosticCommand collects specified container files; parseLinuxDiagnostics preserves availability/validity and parses contention evidence.'),
 ('src/experiment/proxy-config.js', 'sameProxyConfiguration compares canonical equivalent listener/upstream topology; prevents fault injection from reconfiguring an already-correct live proxy.'),
 ('src/experiment/outcomes.js', 'durationSeconds parses lifecycle durations; quantile computes interpolated percentiles; deriveOutcomes joins measured cohort, traces and reconciliation into endpoint/window/fault results.'),
 ('src/experiment/qualification.js', 'qualifyRun combines validity gates, safety, performance and censoring into separate fields and a status; called by qualify-run CLI.'),
 ('src/experiment/analysis.js', 'pairedContrast, holm, decision and restrictedDuration support the planned paired-analysis framework; descriptive paper analysis is a separate entry point.'),
 ('infra/k6/load.js', 'options declares open-loop scheduling; setup makes GET fixture; createTransfer generates mapped payload/key; recordAttempt emits evidence; default selects the mixed operation.'),
 ('infra/k6/smoke.js', 'Small four-iteration HTTP smoke workload, not a controller-managed scientific run.'),
 ('scripts/run-experiment.mjs', 'validateMatrix/select filters rows; k6Arguments constructs load command; runOne operates lifecycle; applyFault/drain/export/qualification/cleanup connect measurement to evidence.'),
 ('scripts/capacity-pilot.mjs', 'assertFrozen and isolation hooks protect a session; runTrial spawns run-experiment with the batch/row; readEvidence feeds protocol evaluation.'),
]


def symbol_locations(path):
    rows = []
    lines = (ROOT / path).read_text().splitlines()
    for n, line in enumerate(lines, 1):
        m = re.match(r'(?:export\s+)?(?:async\s+)?(?:function|class)\s+(\w+)', line)
        if not m:
            m = re.match(r'  (?:async )?(\w+)\(.*\)\s*\{', line)
        if not m:
            m = re.match(r'  (constructor)\(', line)
        if not m and re.match(r'export default function\s*\(', line):
            rows.append(f'default:{n}')
        if m and m.group(1) not in {'if', 'for', 'while', 'switch', 'catch'}:
            rows.append(f'{m.group(1)}:{n}')
    return rows


def file_index():
    parts = []
    for path, description in FILES:
        symbols = symbol_locations(path)
        parts.extend([f'### {Path(path).stem.replace("-", " ")} in {Path(path).parent.as_posix().replace("/", " ")}', '',
                      f'`{path}` — {description}', '',
                      'Source navigation: ' + ('; '.join(f'`{s}`' for s in symbols) if symbols else '`entry point:1`') + '.', ''])
    parts.extend(['Other important entry points are individually traced in section 17. SQL migration responsibilities are in section 10; dashboard/provisioning files in section 14; protocol and test ownership in sections 11 and 16. Historical `continuations/` scripts belong to specific reviewed recovery episodes and are not generic new-pilot launchers.', ''])
    return '\n'.join(parts)


def excerpt(path, symbol):
    lines = (ROOT / path).read_text().splitlines()
    if symbol == 'commitExcerpt':
        start, end = 333, 343
    elif symbol == 'createTransfer':
        start, end = 51, 71
    elif symbol == 'runTrial':
        start, end = 99, 118
    else:
        start = next(i for i, line in enumerate(lines, 1) if re.match(r'  (?:async )?' + symbol + r'\(', line))
        end = next(i for i in range(start + 1, len(lines) + 1) if lines[i - 1] == '  }')
    text = '\n'.join(lines[start - 1:end])
    return f'Source excerpt: `{path}:{start}–{end}`. Original code; visual line wrapping does not change the source.\n\n```javascript\n{text}\n```'


def expanded_markdown(raw):
    # The CLI can print these replacements for apply_patch-based materialization.
    raw = re.sub(r'<!-- CODE ([^ ]+) ([^ ]+) -->', lambda m: excerpt(m[1], m[2]), raw)
    return raw.replace('<!-- FILE_INDEX -->', file_index())


def verify_excerpts(md):
    pattern = r'Source excerpt: `([^`]+):(\d+)–(\d+)`[^\n]*\n\n```javascript\n(.*?)\n```'
    matches = list(re.finditer(pattern, md, re.S))
    if len(matches) != 6:
        raise ValueError(f'Expected six source excerpts, found {len(matches)}')
    for match in matches:
        path, start, end, quoted = match.groups()
        actual = '\n'.join((ROOT / path).read_text().splitlines()[int(start)-1:int(end)])
        if actual != quoted:
            raise ValueError(f'Stale source excerpt at {path}:{start}; refresh documentation before building')


def add_hyperlink(p, text, target=None, anchor=None):
    link = OxmlElement('w:hyperlink')
    if anchor:
        link.set(qn('w:anchor'), anchor)
    else:
        rel = p.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    prop = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '155D85'); prop.append(color)
    run.append(prop)
    t = OxmlElement('w:t'); t.text = text; run.append(t); link.append(run); p._p.append(link)


def inline(p, text, size=None):
    tokens = re.split(r'(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))', text)
    for tok in tokens:
        if not tok:
            continue
        m = re.fullmatch(r'\[([^]]+)\]\(([^)]+)\)', tok)
        if m:
            add_hyperlink(p, m[1], m[2]); continue
        r = p.add_run(tok.strip('`') if tok.startswith('`') else tok.replace('**', ''))
        if size:
            r.font.size = Pt(size)
        if tok.startswith('`'):
            r.font.name = 'Liberation Mono'; r.font.size = Pt(size - .5 if size else 9.5)
        if tok.startswith('**'):
            r.bold = True


def flow_image(kind, out):
    # Compact equivalents of the adjacent Mermaid figures; no external service.
    layouts = {
        0: [
            ['Pilot controller', 'Experiment runner', 'k6 and fault controllers'],
            ['Gateway and async API', 'REST or Kafka workflow', 'Adapter and simulator'],
            ['Application persists and notifies', 'Exported run evidence', 'Qualification and figures'],
        ],
        1: [['RECEIVED / PENDING', 'VALIDATED / PENDING', 'PREPARED / PENDING'],
            ['FULFILLED / COMPLETED or FAILED', 'Commit confirmation', 'Callback and status reads']],
        2: [['callProvider', 'HTTP adapter boundary', 'Adapter service and router'],
            ['toProviderRequest', 'Provider and journal', 'fromProviderResponse'],
            ['Normalized result', 'Provider attempt record', 'Fulfill or fail']],
        3: [['Select row', 'Verified clean start', 'Start and configure stack'],
            ['Warm up', 'Drain reset and register', 'Measure and sample'],
            ['Fault if scheduled', 'Drain and export', 'Derive and qualify'],
            ['Clear faults and network', 'Verify cleanup', 'Record and review']],
    }
    rows = layouts[kind]
    width, bw, bh, gap, margin = 1530, 448, 130, 46, 44
    height = len(rows) * (bh + 72) + 28
    im = Image.new('RGB', (width, height), 'white'); d = ImageDraw.Draw(im)
    font_path = '/System/Library/Fonts/Supplemental/Arial.ttf'
    if not Path(font_path).exists():
        font_path = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    font = ImageFont.truetype(font_path, 30)
    # Snake order avoids crossing connectors and makes the sequence readable.
    previous = None
    for row_index, row in enumerate(rows):
        for index, label in enumerate(row):
            col = index if row_index % 2 == 0 else 2 - index
            x, y = margin + col * (bw + gap), 18 + row_index * (bh + 72)
            if previous:
                px, py, pc, pr = previous
                if pr == row_index:
                    sign = 1 if col > pc else -1
                    start = (px + bw if sign == 1 else px, py + bh // 2)
                    end = (x if sign == 1 else x + bw, y + bh // 2)
                    d.line([start, end], fill='#525D67', width=4)
                    ex, ey = end
                    d.polygon([(ex, ey), (ex - sign * 14, ey - 8), (ex - sign * 14, ey + 8)], fill='#525D67')
                else:
                    cx = x + bw // 2
                    d.line([(cx, py + bh), (cx, y)], fill='#525D67', width=4)
                    d.polygon([(cx, y), (cx - 8, y - 14), (cx + 8, y - 14)], fill='#525D67')
            d.rounded_rectangle([x, y, x + bw, y + bh], radius=10, fill='#F1F5F8', outline='#71818D', width=2)
            words = textwrap.wrap(label, 26)
            for j, word in enumerate(words):
                box = d.textbbox((0, 0), word, font=font)
                d.text((x + (bw - (box[2] - box[0])) / 2, y + (bh - 37 * len(words)) / 2 + 37*j), word, font=font, fill='black')
            previous = (x, y, col, row_index)
    im.save(out)


def render(md):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.72)
    sec.left_margin = sec.right_margin = Inches(.8)
    sec.header_distance = sec.footer_distance = Inches(.3)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        s = doc.styles[name]; s.font.name = 'Arial'; s.font.color.rgb = RGBColor(0, 0, 0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    for name, size, before, after in [('Title', 25, 0, 10), ('Heading 1', 17, 15, 8), ('Heading 2', 13, 12, 6), ('Heading 3', 11, 9, 5)]:
        s = doc.styles[name]; s.font.size = Pt(size); s.font.bold = name != 'Title'
        s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    from docx.enum.style import WD_STYLE_TYPE
    code = doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = 'Liberation Mono'; code.font.size = Pt(8.5)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.left_indent = Inches(.12)
    code.paragraph_format.widow_control = False
    header = sec.header.paragraphs[0]
    header.text = 'MOBILE MONEY COORDINATION EXPERIMENT  |  TECHNICAL GUIDE'
    header.runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('21 September 2026  •  '); r.font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    doc.core_properties.title = 'Mobile Money Coordination Experiment Technical Guide'
    doc.core_properties.subject = 'Codebase onboarding and operational reference'
    doc.core_properties.author = 'Project documentation'

    headings = re.findall(r'^## (.+)$', md, re.M)
    bookmark_names = {h: f'section_{i+1}' for i, h in enumerate(headings)}
    lines = md.splitlines(); i = 0; diagram = 0; made_toc = False
    assets = tempfile.TemporaryDirectory(prefix='coordination-doc-diagrams-')
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.startswith('<!--'):
            i += 1; continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title'); i += 1; continue
        if line.startswith('## '):
            if not made_toc:
                doc.add_paragraph('Contents', 'Heading 1')
                for h in headings:
                    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
                    add_hyperlink(p, h, anchor=bookmark_names[h])
                doc.add_page_break(); made_toc = True
            title = line[3:]
            p = doc.add_paragraph(title, 'Heading 1')
            mark = OxmlElement('w:bookmarkStart'); mark.set(qn('w:id'), str(headings.index(title)+1)); mark.set(qn('w:name'), bookmark_names[title])
            p._p.insert(0, mark)
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(headings.index(title)+1)); p._p.append(end)
            i += 1; continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2'); i += 1; continue
        if line.startswith('```'):
            language = line[3:]; block = []; i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i]); i += 1
            i += 1
            if language == 'mermaid':
                out = Path(assets.name) / f'diagram-{diagram}.png'
                flow_image(diagram, out); diagram += 1
                p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(out), width=Inches(6.85))
                caption = ['Logical execution and evidence sequence, not direct network connections. The adapter returns to the application for persistence and callbacks. Kafka additionally uses Redpanda.',
                           'Transfer lifecycle. Failure is permitted from nonterminal states; successful fulfillment precedes commit confirmation.',
                           'Common adapter path. Both coordinators use the same boundary and profile normalization.',
                           'Managed run order. Warm-up and measurement are separate k6 invocations.'][diagram-1]
                p = doc.add_paragraph(caption); p.runs[0].italic = True; p.runs[0].font.size = Pt(9)
            else:
                # Let Word soft-wrap long lines; do not insert newlines inside
                # JSON strings or commands that readers may copy from the guide.
                for n, row in enumerate(block):
                    p = doc.add_paragraph(row, 'Source Code')
                    p.paragraph_format.keep_with_next = n < len(block)-1 and len(block) <= 26
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+', c) for c in cells):
                    rows.append(cells)
                i += 1
            ncols = len(rows[0]); table = doc.add_table(rows=0, cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
            widths = [2.25, 4.65] if ncols == 2 else [2.15, 2.27, 2.48]
            for c, w in zip(table.columns, widths): c.width = Inches(w)
            pr = table._tbl.tblPr
            borders = OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el = OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            for ri, values in enumerate(rows):
                row = table.add_row()
                trpr = row._tr.get_or_add_trPr()
                cant = OxmlElement('w:cantSplit'); trpr.append(cant)
                if ri == 0:
                    repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
                for ci, (cell, value) in enumerate(zip(row.cells, values)):
                    cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr = cell._tc.get_or_add_tcPr()
                    mar = OxmlElement('w:tcMar')
                    for edge, val in [('top','90'),('bottom','90'),('left','95'),('right','95')]:
                        e=OxmlElement('w:'+edge); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); mar.append(e)
                    tcpr.append(mar)
                    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.03
                    inline(p, value, 10)
                    if ri == 0:
                        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E2EBF1'); tcpr.append(shade)
                        for r in p.runs: r.bold=True
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        p = doc.add_paragraph(); inline(p,line)
        if line.startswith('Source excerpt:'):
            p.paragraph_format.keep_with_next=True
            for r in p.runs: r.font.size=Pt(9)
        i += 1
    doc.save(OUTPUT)
    assets.cleanup()
    print(OUTPUT)


if __name__ == '__main__':
    import sys
    raw = SOURCE.read_text()
    if '--replacements' in sys.argv:
        import json
        replacements = {m[0]: excerpt(m[1],m[2]) for m in re.finditer(r'<!-- CODE ([^ ]+) ([^ ]+) -->', raw)}
        if '<!-- FILE_INDEX -->' in raw:
            replacements['<!-- FILE_INDEX -->'] = file_index()
        print(json.dumps(replacements))
    else:
        expanded = expanded_markdown(raw)
        verify_excerpts(expanded)
        render(expanded)
