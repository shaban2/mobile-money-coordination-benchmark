"""Revise a COPY of the supplied pre-results paper using audited local results.

Usage: python3 scripts/build-exploratory-paper.py /path/to/original.docx
Requires python-docx 1.2.0. Run analysis and figure scripts first.
The original file and all experiment evidence are read-only inputs.
"""
import hashlib
import json
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / '.research' / 'exploratory-paper-v1'
original = Path(sys.argv[1]).resolve()
target = ROOT / 'docs' / 'REST_Kafka_Exploratory_Findings_Draft.docx'
assert original != target
source_hash = hashlib.sha256(original.read_bytes()).hexdigest()
data = json.loads((ART / 'analysis.json').read_text())
runs = {(r['key'], r['architecture']): r for r in data['runs']}
sr, sk = runs['screening', 'REST'], runs['screening', 'Kafka']
ar, ak = runs['adapter', 'REST'], runs['adapter', 'Kafka']
dr, dk = runs['database', 'REST'], runs['database', 'Kafka']
archived = data['archivedNoFaultRuns']
# Supplementary pairs whose application-defining files match the reporting set, allowing
# at most the k6 check-threshold comparison operator (calibration used '>' rather than '>=').
same_app = [r for r in archived if set(r['applicationFilesDiffering']) <= {'infra/k6/load.js'}]
other_app = [r for r in archived if r not in same_app]
same_pairs = sorted({r['key'] for r in same_app}, key=lambda k: [r['key'] for r in archived].index(k))
kafka_first_pairs = [k for k in same_pairs if next(r for r in same_app if r['key'] == k and r['firstInPair'])['architecture'] == 'Kafka']
def arch_runs(rs, a): return [r for r in rs if r['architecture'] == a]
def rng(vals, fmt='{:.0f}'):
    lo, hi = min(vals), max(vals)
    return fmt.format(lo) if lo == hi else f'{fmt.format(lo)}–{fmt.format(hi)}'
same_r, same_k = arch_runs(same_app, 'REST'), arch_runs(same_app, 'Kafka')
assert len(same_pairs) == 3 and len(kafka_first_pairs) == 2 and len(other_app) == 2
words = {1: 'one', 2: 'two', 3: 'three', 4: 'four'}
# Public artifact: the redacted derivative repository and its Zenodo archive (DOI filled in once minted).
ARTIFACT_URL = 'https://github.com/shaban2/mobile-money-coordination-benchmark'
ARTIFACT_VERSION = 'v1.2.0'
ARTIFACT_DOI = None
artifact_ref = f' Zenodo, 2026. doi: {ARTIFACT_DOI}.' if ARTIFACT_DOI else ' GitHub, 2026.'
artifact_cite = (f'archived at Zenodo (doi: {ARTIFACT_DOI})' if ARTIFACT_DOI else 'with an archival identifier pending')
fresh = {(r['key'], r['architecture']): r for r in data['freshFaultRuns']}
far, fak = fresh['adapter', 'REST'], fresh['adapter', 'Kafka']
fdr, fdk = fresh['database', 'REST'], fresh['database', 'Kafka']
assert fak['firstInPair'] and fdk['firstInPair'] and len(data['freshFaultRuns']) == 4
def infault(r): return r['faultInterval']['completionsDuringFault']
def compare(a, b, what=None):
    if a == b: return 'equal for REST and Kafka'
    return f'{"higher" if a > b else "lower"} for REST ({a}) than for Kafka ({b})'
def date_of(r): return r['measurementStart'][:10]
doc = Document(original)
# Retain the original title/author block, its continuous column break, and final
# section properties. The approved change of study scope requires a body rewrite.
body = doc._element.body
first_four = list(body)[:4]
for child in list(body):
    if child not in first_four and child.tag != qn('w:sectPr'):
        body.remove(child)
title = 'REST Orchestration or Kafka Choreography Behind a Unified Telecom Mobile-Money P2P Transfer API: A Controlled Synthetic Benchmark'
doc.paragraphs[0].text = title
doc.paragraphs[0].style = doc.styles['Title']
for name in ['Title', 'Heading 1', 'Heading 2', 'Caption']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    for el in list(style.element.iter()):
        if el.tag in {qn('w:pBdr'), qn('w:spacing')} and el.getparent().tag == qn('w:rPr'):
            el.getparent().remove(el)
    for el in list(style.element.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
    fonts = style.element.find('.//' + qn('w:rFonts'))
    if fonts is not None:
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            fonts.attrib.pop(qn('w:' + attr), None)
for r in doc.paragraphs[0].runs:
    r.font.name = 'Times New Roman'; r.font.size = Pt(22); r.font.bold = False
    r.font.color.rgb = RGBColor(0, 0, 0)
doc.paragraphs[0].paragraph_format.space_after = Pt(8)

manuscript = [f'# {title}', '', 'Shaban Lubanga', '', doc.paragraphs[2].text, '']

def para(text, style=None, size=10, indent=True, italic=False, bold=False):
    p = doc.add_paragraph(style=style)
    f = p.paragraph_format
    f.space_before = Pt(0); f.space_after = Pt(0); f.line_spacing = 1
    f.first_line_indent = Inches(.14) if indent else Inches(0)
    f.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text); r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0); r.italic = italic; r.bold = bold
    manuscript.extend([text, ''])
    return p

def heading(text, sub=False):
    p = para(text, 'Heading 2' if sub else 'Heading 1', indent=False, italic=sub)
    p.paragraph_format.space_before = Pt(5 if sub else 7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if sub else WD_ALIGN_PARAGRAPH.CENTER
    manuscript[-2] = ('### ' if sub else '## ') + text

def figure(name, caption, alt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ART / 'figures' / f'{name}.png'), width=Inches(3.32))
    shape._inline.docPr.set('descr', alt)
    cap = para(caption, size=8, indent=False)
    cap.paragraph_format.space_after = Pt(6)
    cap.paragraph_format.keep_together = True
    manuscript.insert(len(manuscript)-2, f'![{alt}](figures/{name}.png)\n')

def table(title, headers, rows, widths):
    p = para(title, size=8, indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
    for c, w in zip(t.columns, widths): c.width = Inches(w)
    for i, row in enumerate([headers] + rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[j]); cell.text = str(value)
            margins = OxmlElement('w:tcMar')
            for side in ['left', 'right']:
                item = OxmlElement('w:' + side); item.set(qn('w:w'), '45'); item.set(qn('w:type'), 'dxa'); margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            cp = cell.paragraphs[0]; cp.paragraph_format.space_after = Pt(3)
            cp.paragraph_format.space_before = Pt(3); cp.paragraph_format.line_spacing = 1
            cp.paragraph_format.keep_with_next = i < len(rows)
            cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            for r in cp.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(8); r.bold = i == 0
            if i == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'ECEFF1'); cell._tc.get_or_add_tcPr().append(shd)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'bottom', 'insideH']:
        e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'B8BDC1'); borders.append(e)
    t._tbl.tblPr.append(borders)
    manuscript.extend(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---']*len(headers)) + ' |'])
    manuscript.extend('| ' + ' | '.join(map(str, row)) + ' |' for row in rows)
    manuscript.append('')

def ms_display(value):
    # Remove sub-nanosecond binary interpolation noise before decimal display.
    return format(Decimal(str(round(value, 8))).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP), ',.1f')

para(f'Abstract. We report an exploratory comparison of REST-style orchestration and Kafka-based choreography for synthetic transfers under a common asynchronous HTTP contract. Both implementations use co-located workflow logic, a shared PostgreSQL ledger design, and the same provider simulator; Kafka additionally uses a Redpanda broker and durable inbox/outbox processing. The reporting set comprises six runs: one REST–Kafka pair at 8 operations/s without faults and two pairs at 4 operations/s with an adapter outage or database-response delay; two further Kafka-first fault pairs and six archived no-fault runs are reported beside it. All {data["measuredCompletions"]:,} measured transfers in the reporting set, and all {sum(r["counts"]["completed"] for r in data["freshFaultRuns"]):,} in the fresh fault runs, completed correctly, with no observed ledger-invariant violations or drain-only completions. At 8 operations/s, observed p95 ingress-to-commit-acknowledgement latency was 59.0 ms for REST and 172.1 ms for Kafka. Three archived 4 operations/s no-fault pairs with identical application files, two of them executed Kafka first, showed the same direction and size of difference. Faults produced temporary backlogs and longer completion tails in both implementations; during the database-response delay REST kept completing at about 1.2 transfers/s per window while Kafka fell to about 0.25, consistent with Kafka issuing about 1.7 times as many database statements per run. Two fresh fault pairs collected afterwards with the Kafka condition executed first gave in-fault completions of {infault(fdr)} for REST and {infault(fdk)} for Kafka under database delay and {infault(far)} and {infault(fak)} under the adapter outage, with pre-fault p95 of {fdr["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms and {fdk["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms. REST also had lower sampled service-set CPU and memory use. These are configuration-specific observations from one execution per architecture per condition and order, not capacity estimates or statistical evidence of general superiority. We provide explicit accounting boundaries and reproducible local analysis artifacts, and identify further run-level replication and configuration sensitivity as the main remaining uncertainties.', size=9, indent=False, italic=True)
para('Index Terms. Synthetic benchmark, transfer coordination, Kafka, orchestration, choreography, durable completion.', size=9, indent=False)

heading('I INTRODUCTION')
para('Mobile money lets people store and move value through an account tied to a mobile phone number, usually operated by a telecom provider. Because a sender and a recipient may use different providers, transfers between them depend on interoperability, which research on mobile payments identifies as a continuing economic and policy problem [13]; payment-policy bodies recommend harmonised application programming interfaces (APIs) because fragmented API standards raise processing time, cost, and error risk [14].')
para('A person-to-person (P2P) transfer is one person sending money from their mobile wallet to another person\'s wallet, much like handing over cash, except that the two wallets may belong to different providers. Fig. 1 shows such a transfer as the user sees it. The sender chooses a recipient and an amount in a wallet at Provider A, the transfer application submits one request to a unified P2P transfer API, the platform routes the transfer to the recipient\'s provider and records it, and the recipient at Provider B receives the value. The sender then sees COMPLETED or FAILED; when processing takes longer, the application first shows PENDING and obtains the final result later through a callback or a status check. Behind that simple picture the platform must translate different provider interfaces, prevent duplicate transfers, keep balances correct, recover from failures, and reliably deliver the final result.')
figure('user-view', 'Fig. 1. User view of a P2P transfer through a unified mobile-money API. Adapted from the author\'s thesis proposal, based on the GSMA P2P use case [1]. The shaded step is what this paper measures; provider-side processing is simulated.', 'Four steps: the sender at Provider A chooses recipient and amount, a transfer application submits one request, the unified API routes and processes the transfer, and the recipient at Provider B receives the value; the sender receives COMPLETED or FAILED, possibly after PENDING.')
para('The GSMA Mobile Money API specification [15] defines the message language (party identifiers, amount, currency, transaction references, and status) and its asynchronous callback and polling flows; its P2P use case [1] frames the transfer in Fig. 1. Mojaloop [16] helps explain provider boundaries and a staged lookup, quote, and transfer workflow. Neither decides how the platform coordinates its own work. This paper varies one internal design choice while everything else stays fixed: REST orchestration, where a coordinator calls the workflow steps in sequence, versus Kafka choreography, where co-located handlers react to events carried through a broker and durable database records. Both implementations acknowledge the request first (HTTP 202) and deliver the result later. Because that early acknowledgement can arrive before the transfer is safely stored, we measure each transfer from the moment the API receives it to the moment the database confirms its final commit, and we count only transfers that completed, were not duplicated, and left the ledger balanced.')
para('The prototype is deliberately bounded. Each transfer is routed through an adapter to one of two synthetic provider interfaces and recorded in one local ledger; it models the provider-facing side of such a platform, not settlement between two independently operated ledgers, and no real money or customer data is involved.')
para('Our contribution is an auditable comparison that links generated operations to unique durable completions, ledger checks, and fault-response trajectories. The exploratory scope and the descriptive analysis were chosen after observing the results. We do not claim preregistration, a new coordination pattern, or a universal architecture ranking.')
heading('Research questions', True)
para('RQ1. At the tested no-fault load, what correct goodput and durable-completion latency were observed for each implementation?', indent=False)
para('RQ2. At a fixed 4 operations/s, how did an adapter outage and a database slowdown affect completions during the fault, the backlog, and the return to normal?', indent=False)
para('RQ3. Did every transfer end with a correct ledger record, and how much CPU and memory did each implementation use while doing so?', indent=False)

heading('II RELATED WORK')
para('Kazanavičius and Mažeika compare HTTP REST, Kafka, RabbitMQ, gRPC, and GraphQL in communication experiments for monolith decomposition [2]. Their comparison motivates careful attention to the workload and deployment when interpreting transport results. Our endpoint instead requires a reconciled transfer and observed terminal database acknowledgement, and includes deliberately injected dependency faults.')
para('Financial-style transfer benchmarks already exist. Son, Lee, and Lee describe a transfer-process model comparing two-phase commit, Saga choreography, Saga orchestration, and Axon [3]. The accessible author abstract supports this domain overlap; we do not infer its replication schedule or claim equivalent measurements without the full methods. Our contribution is therefore not the first use of transfers to compare coordination approaches.')
para('Kristianto and Zahra vary service count, concurrent users, and instances in distributed Saga implementations [4]. Both of their coordination models use Kafka, unlike our direct-call versus event-driven packages. Their reported choreography advantage is tied to that deployment and does not contradict a different observation in our co-located implementation. We do not compare numerical latencies across these studies.')
para('Nadeem and Malik report initial findings from porting a benchmark system to a workflow engine and identify fuller evaluation as future work [9]; we adopt the same early-results positioning. Hegde et al. compare synchronous and asynchronous microservice implementations using k6, Docker resource monitoring, and worker termination, and report API-response performance [10]. Our endpoint is durable ledger completion rather than API response, and both of our conditions expose an asynchronous client contract. Aydın and Çebi examine choreography and orchestration Saga implementations with isolated databases and cross-service rollback [11]; our common ledger is not that setting. The inbox/outbox mechanism in our Kafka condition follows the transactional outbox pattern described by Richardson [12].')
para('Hasselbring argues that useful benchmarking requires explicit workloads, fairness, verifiability, and a replication package [5]. Kalibera and Jones emphasize repetition at the level where performance varies [6]. These considerations motivate our documented scope and our refusal to treat thousands of transfers within a run as thousands of independent architecture experiments.')

heading('III SYSTEM UNDER TEST')
figure('architecture', 'Fig. 2. Deployment boundary. REST and Kafka are alternative application conditions; internal workflow stages are not independently deployed services.', 'Two alternative application processes share the same types of provider, ledger and callback services. Only the Kafka condition uses an external broker.')
heading('Common contract and provider simulation', True)
para('POST /v1/transfers accepts payerId, payeeId, amount, currency, clientReference, a provider profile, and an optional callback URL. An idempotency key identifies the operation. New accepted work receives HTTP 202; replays refer to the original transfer. GET /v1/transfers/{id} returns its public state. Provider adapters translate the canonical fields and responses to two synthetic interfaces. No production provider data or real money is used.')
para('Internal states progress through RECEIVED, VALIDATED, and PREPARED to FULFILLED or FAILED. The public states are PENDING, COMPLETED, and FAILED. Successful completion creates balanced local debit/credit postings. A durable synthetic-provider journal supports idempotent effects across adapter restarts. This is not a conformance test of any standard and not proof of atomic settlement across a provider and our ledger.')
heading('Implementation packages', True)
para('REST orchestration invokes local workflow functions in sequence and uses the common HTTP provider boundary. It does not make a REST network call between every internal stage. Kafka choreography subscribes co-located handlers to workflow events. PostgreSQL inbox/outbox records couple handler progress to event production and duplicate detection; a dispatcher polls every 50 ms and also dispatches on selected execution paths. KafkaJS communicates with one Redpanda broker using four partitions per workflow topic and replication factor one.')
para('The business contract, provider delay, retry configuration, and four-worker limit are shared. Kafka-specific database writes, broker work, and dispatch scheduling are part of its measured implementation. Consequently, this comparison changes a package of coordination and durability mechanisms; it does not isolate the causal cost of the Kafka transport alone.')
heading('Environment and controls', True)
para('All runs used one Apple M2 host with eight logical CPUs and 16 GiB physical memory, running Linux containers through Docker. The application limit was two CPUs and 1 GiB memory. Software included Node.js 24 container images, pg 8.22.0, KafkaJS 2.2.4, PostgreSQL 17, Redpanda v26.1.14, Toxiproxy 2.12.0, and k6 1.8.0. Image IDs, resolved configuration, host metadata, and source archives were retained; tags alone are not the identity record.')
para('The provider simulator used a 25 ms configured delay. Both conditions used eight provider retry attempts with 5 s spacing and a 90 s workflow timeout. The standard gateway profile applied 10 ms delay, 2.5 ms jitter, and 0.05% loss on each gateway egress. This profile was unchanged between conditions, although the primary latency clock begins after the request reaches the application. Monitoring remained enabled, and unrelated containers were paused during collection.')

heading('IV EXPERIMENTAL METHOD')
heading('Dataset and study history', True)
para(f'The reporting set contains all six runs in the separately approved higher-load and fixed-load fault sequence, completed on 21 September 2026 under one frozen source and image identity (Section VII). Table I states the actual durations. Earlier engineering checks, calibration runs, interrupted attempts, and runs under previous harness revisions remain archived. Three archived no-fault pairs at 4 operations/s were collected under earlier experiment-harness revisions but identical application files; Section V reports them as separately labelled supplementary observations, not pooled with the reporting set. After the reporting set was analysed, two further fault pairs were collected on {date_of(fak)} with the Kafka condition seeded to run first, under a rebuilt implementation whose application files differ from the reporting set only in the package name (source identity {data["freshFaultSourceSha256"][:12]}…); Section V reports them beside the original pairs. The original capacity-search and provisional 96-run confirmatory plans were not completed.')
table('TABLE I  COMPLETED EXPLORATORY DESIGN', ['Condition', 'Rate\n(operations/s)', 'Warm-up\n(min)', 'Measure\n(min)', 'Runs'], [
    ['No fault', '8', '2', '4', 'R + K'], ['Adapter outage', '4', '5', '10', 'R + K'], ['Database delay', '4', '5', '10', 'R + K']], [.92, .78, .56, .56, .50])
para('Each row has one REST execution and one Kafka execution. The seeded order generator selected REST then Kafka in all three reporting-set pairs; the realized order was not balanced. Within the reporting set we therefore cannot separate an order or host-time effect from architecture. The two fresh fault pairs were executed Kafka first and provide the order-balanced counterpart for the fault scenarios; the supplementary archived pairs do so for the no-fault condition. No capacity boundary was established, and the fault rate was an absolute 4 operations/s, not 90% of estimated capacity.')
figure('method', 'Fig. 3. Experimental method. Every run follows the same frozen pipeline; the three conditions feed the three research questions through the tables and figures named.', 'Three bands: the three test conditions, the six-step pipeline every run follows, and the measures that answer each research question.')
heading('Deterministic workload', True)
para('k6 uses the constant-arrival-rate executor to schedule operations independently of completion time while virtual users are available [7]. An operation is either a creation attempt or a status read, not necessarily a new transfer. The repeating 100-operation schedule contains 76 unique creations, four idempotent creation replays, and 20 status reads. Thus replays are 5% of creation attempts and 4% of all operations; these proportions are benchmark choices rather than estimates from customer traffic.')
para('The iteration index generates amounts as 100 + (sourceIteration mod 900), encoded as positive integer minor-unit strings with currency UGX. A seed-derived offset rotates numbered synthetic payer/payee identities; iteration parity selects the provider profile. Replay slots reuse the preceding creation payload and key. Every status read targets the same setup transfer, which is excluded from the measured cohort; this is a deliberately simple read workload, not realistic polling.')
para('The scheduled counts are 1,920 operations for a 4-minute run at 8/s and 2,400 for a 10-minute run at 4/s. Recorded streams contain one extra operation in some runs because an operation may fall on the ending boundary. A transfer belongs to the measured cohort if the server received it inside the measurement interval; setup transfers and boundary operations are retained in the accounting but excluded from the cohort.')
heading('Procedure and fault injection', True)
para('The controller prepares the selected condition, verifies configuration and isolation, runs warm-up, and waits for pending work to drain. It then resets trial data before starting the measured k6 phase. Fault timing uses the k6 scenario-start marker. After measurement it allows up to 10 minutes for drain, exports evidence, evaluates validity and outcomes, and cleans up; these steps are additional to the durations in Table I.')
para('At measurement second 240, the adapter scenario stops the common adapter container, waits 30 s, and restarts it. The observed outage boundary runs from stop acknowledgement to health readiness, about 30.44 s. This is controlled adapter unavailability, not an abrupt application or broker crash. The database scenario adds 100 ms of downstream PostgreSQL-response delay through Toxiproxy for about 60.12 s. It does not slow disk hardware or isolate database-server execution time. Load continues at 4 operations/s throughout both faults and recovery.')
heading('What is measured and how a run is accepted', True)
para('For each unique transfer, latency is application API ingress to application-observed acknowledgement of the terminal database COMMIT, using one application wall clock. The clock starts before the request body is read and stops after COMMIT returns; the value is then persisted separately. It excludes client-to-API transit and is not the database’s exact internal commit instant. HTTP 202, an internal state timestamp, and callback delivery are not substitutes.')
para('Correct goodput counts reconciled unique successes both admitted and confirmed during measurement, divided by measurement seconds. Latency quantiles describe those successful completions. We separately retain failures, pending transfers, missing timing, and completions during drain. Backlog is the number of measured arrivals still without a commit confirmation at each 30 s boundary; it is a sample, not the maximum queue size.')
para('Recovery is the time from observed fault clearance to the end of three consecutive complete 30 s windows under continued load, each reaching at least 90% of baseline goodput and no more than 110% of baseline p95 latency. Baseline uses the last three full pre-fault windows. Backlog clearance is the first full post-clearance window ending at or below the baseline maximum backlog. Both are descriptive, and because a candidate window must begin after clearance, the rule adds a fixed reporting delay. We therefore also report the window-level trajectory: completions in windows whose boundary fell inside the fault interval, the first window after clearance, and whether the first complete window entirely after clearance already met the baseline criteria. That description was chosen after seeing the trajectories.')
para('Validity checks cover configuration identity, workload delivery, correlation, instrumentation and service coverage, fault evidence, and cleanup. Correctness checks include unique idempotency keys, terminal-state history, one balanced ledger transaction per fulfilled transfer, posting/amount agreement, account-balance reconstruction, and conservation by currency. Final callback identity and status are checked separately. Valid performance failures and censored observations are retained rather than treated as invalid instrumentation.')
heading('Descriptive analysis and resource scope', True)
para('We recompute archived outcomes from traces, snapshots, and invariants and require exact agreement before generating tables. Quantiles use linear interpolation at (n − 1)p. With one pair per condition and order, we report no confidence interval, significance test, or equivalence conclusion, and conditions with different rates and durations are not pooled. For the supplementary archived pairs we verify qualification, invariant checks, and the p95 recomputation, and compare every application-defining file (application source outside the experiment harness, migrations, contracts, container definitions, k6 script, and dependency lock) in each frozen source archive against the reporting-set archive.')
para('Resource summaries sum the application, gateway, PostgreSQL, Toxiproxy, adapter, callback receiver, Prometheus, and Grafana, plus Redpanda for Kafka. They exclude k6, the controller, and host overhead. Memory is Docker CLI usage, which subtracts cache on Linux [8]; CPU sums container percentages divided by 100. Means are time-weighted over samples about 2 s apart within measurement; a stopped adapter contributes zero while absent. These are descriptive footprints, not cost or reserved capacity.')

heading('V RESULTS')
para('Reading the measures. Completion latency is the time from the API receiving a request to the database confirming its final commit, in milliseconds (1,000 ms = 1 s). The median is the time within which half of the transfers completed; p95 is the time within which 95 of every 100 completed, and p99 the time within which 99 of every 100 completed, so p95 and p99 describe the slow tail. Correct goodput is the number of completed, non-duplicated, balanced transfers per second. Backlog is the number of accepted transfers not yet committed at a 30-second boundary. CPU is reported in core equivalents: 0.25 means a quarter of one processor core busy on average.', indent=False)
heading('Completion and correctness', True)
para(f'All six runs were qualified. Across the reporting set, {data["measuredCompletions"]:,} unique measured transfers completed correctly. Each no-fault run completed {sr["counts"]["completed"]:,} transfers, and each fault run completed {ar["counts"]["completed"]:,}. Every measured transfer had usable timing and a matching final callback. There were zero measured failures, zero pending transfers after drain, zero drain-only completions, and no observed invariant violations.')
labels = {'screening': 'No fault', 'adapter': 'Adapter', 'database': 'Database'}
table('TABLE II  OBSERVED COMPLETIONS AND LATENCY', ['Condition', 'Completed\ntransfers', 'Correct goodput\n(transfers/s)', 'p95 latency\n(ms)', 'p99 latency\n(ms)'], [
    [labels[r['key']] + ' ' + ('R' if r['architecture'] == 'REST' else 'K'), f'{r["counts"]["completed"]:,}', f'{r["goodput"]:.3f}', ms_display(r['latency']['p95']), ms_display(r['latency']['p99'])] for r in data['runs']], [.80, .58, .70, .62, .62])
para('R denotes REST and K denotes Kafka. Completed transfers exclude replays and status reads. Latency is the time from API ingress to the confirmed final commit; p95 means 95 of every 100 transfers completed within that time, so 2,479.4 ms is about 2.5 s. Equal correct goodput follows from every implementation completing the whole offered cohort; it does not indicate equal capacity.', size=8, indent=False)
heading('No-fault observation', True)
para(f'At 8 operations/s, both implementations achieved {sr["goodput"]:.3f} correct transfers/s and had zero backlog at all eight 30 s boundaries. Median completion latency was {sr["latency"]["median"]:.0f} ms for REST and {sk["latency"]["median"]:.0f} ms for Kafka. The p95 (95 of every 100 transfers) was {sr["latency"]["p95"]:.1f} ms for REST and {sk["latency"]["p95"]:.1f} ms for Kafka, an observed difference of {sk["latency"]["p95"]-sr["latency"]["p95"]:.1f} ms. Fig. 4 shows the within-run distributions. A four-minute screen at one rate does not demonstrate long-run sustainability or maximum capacity.')
figure('screening-cdf', 'Fig. 4. Empirical completion-latency distributions at 8 operations/s without faults. Each curve is one run with 1,459 measured completions, not a distribution of independent run estimates.', 'REST and Kafka empirical completion-latency distributions, each from one no-fault execution at 8 operations per second.')
heading('Fault response', True)
para(f'The adapter outage produced a sampled backlog maximum of {ar["maxSampledBacklog"]} in each implementation (Fig. 5). Over the whole measurement, 95 of every 100 transfers completed within {ar["latency"]["p95"]/1000:.2f} s for REST and {ak["latency"]["p95"]/1000:.2f} s for Kafka; the p99 reached {ar["latency"]["p99"]/1000:.2f} and {ak["latency"]["p99"]/1000:.2f} s. Both eventually completed every measured transfer before measurement ended. Whole-run goodput of 3.04/s masks the temporary interruption visible in the trajectory. In the one window whose boundary fell inside the outage, each implementation completed {ar["faultInterval"]["completionsDuringFault"]} transfer; in the following window each completed the backlog at {ar["faultInterval"]["catchUpWindow"]["correctGoodput"]:.2f} transfers/s. At this resolution the two adapter-outage trajectories are indistinguishable except for the latency offset already present before the fault.')
figure('adapter-timeline', 'Fig. 5. Adapter-outage response at 4 operations/s. Shading spans the observed fault intervals. Horizontal segments show statistics for complete 30 s windows; backlog markers show boundary samples. Lines between backlog samples are visual guides.', 'Goodput, sampled backlog and window p95 latency during the adapter outage and continued load.')
para(f'Under database-response delay, sampled backlog reached {dr["maxSampledBacklog"]} for REST and {dk["maxSampledBacklog"]} for Kafka (Fig. 6). Measurement-wide p95 was {dr["latency"]["p95"]/1000:.2f} versus {dk["latency"]["p95"]/1000:.2f} s. Both resumed sufficient completion activity to clear all measured work before the end. The two implementations behaved differently while the delay was active: in the two windows whose boundaries fell inside it, REST completed {dr["faultInterval"]["completionsDuringFault"]} transfers (about {dr["faultInterval"]["meanGoodputDuringFault"]:.2f}/s per window) while Kafka completed {dk["faultInterval"]["completionsDuringFault"]} (about {dk["faultInterval"]["meanGoodputDuringFault"]:.2f}/s). The catch-up window then reached {dr["faultInterval"]["catchUpWindow"]["correctGoodput"]:.2f}/s for REST and {dk["faultInterval"]["catchUpWindow"]["correctGoodput"]:.2f}/s for Kafka. The longest single completion was {dr["latency"]["max"]/1000:.1f} s for REST and {dk["latency"]["max"]/1000:.1f} s for Kafka against the 90 s workflow timeout, so a longer delay of this kind would have produced timeouts in Kafka first. These observations do not estimate an isolated fault effect relative to the 8/s screen because load and duration differ; the within-run pre-fault baseline and the supplementary 4/s pairs provide the local reference.')
figure('database-timeline', 'Fig. 6. Database-delay response at 4 operations/s, with the same display conventions as Fig. 5. The injected delay affects downstream database responses for about 60 s.', 'Goodput, backlog and window p95 latency during the database-response delay; Kafka has a larger sampled backlog in this execution.')
heading('Fresh Kafka-first fault pairs', True)
para(f'Two fresh pairs repeated the adapter-outage and database-delay scenarios at 4 operations/s with identical settings, durations and workload seeds, but with the Kafka condition executed first. All four runs qualified with zero failures, zero pending transfers, and no invariant violations; every measured transfer had a matching final callback. Table VI places each fresh pair beside the original REST-first pair. Pre-fault baseline p95 was {far["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms and {fdr["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms for REST and {fak["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms and {fdk["fault"]["baseline"]["p95DurableCompletionMs"]:.0f} ms for Kafka, in the range of every earlier 4 operations/s observation.')
para(f'Under the adapter outage the fresh pair completed {infault(far)} (REST) and {infault(fak)} (Kafka) transfers in the window ending inside the outage, against {infault(ar)} and {infault(ak)} originally; sampled backlog peaked at {far["maxSampledBacklog"]} and {fak["maxSampledBacklog"]}. Under database delay the fresh pair completed {infault(fdr)} (REST) and {infault(fdk)} (Kafka) transfers in the two windows ending inside the delay, against {infault(dr)} and {infault(dk)} originally, with peak backlog {fdr["maxSampledBacklog"]} and {fdk["maxSampledBacklog"]} and longest completions of {fdr["latency"]["max"]/1000:.1f} s and {fdk["latency"]["max"]/1000:.1f} s. With one pair per order, these are two observations per scenario, not an estimate of order effect; they let the direction seen in the original pairs be checked against a pair in which Kafka did not run second.')
def order_rows(orig_r, orig_k, fr, fk, label):
    rows = []
    for tag, r in [('orig.', orig_r), ('orig.', orig_k), ('fresh', fk), ('fresh', fr)]:
        rows.append([f'{label} {tag}', '1st' if r['firstInPair'] else '2nd', 'R' if r['architecture'] == 'REST' else 'K', f"{r['counts']['completed']:,}",
                     r['maxSampledBacklog'], infault(r), f"{r['latency']['p95']/1000:.2f}", 'yes' if r['faultInterval']['firstWindowFullyAfterClearanceMeetsBaseline'] else 'no'])
    return rows
table('TABLE VI  ORIGINAL AND FRESH 4 OPS/S FAULT PAIRS BY REALIZED ORDER', ['Fault / pair', 'Order', 'Arch.', 'n', 'Peak\nbacklog', 'Done in\nfault', 'p95\n(s)', 'Normal\n1st win.'],
    order_rows(ar, ak, far, fak, 'Adapter') + order_rows(dr, dk, fdr, fdk, 'Database'), [.66, .36, .34, .40, .40, .40, .34, .40])
para(f'Original pairs ran REST first on 21 September 2026 under the reporting-set identity; fresh pairs ran Kafka first on {date_of(fak)} under the rebuilt identity (Section VII). Columns are defined as in Table III. Measurement-wide p95 mixes pre-fault, in-fault and catch-up completions and is shown for comparability with Table II only.', size=8, indent=False)
table('TABLE III  DESCRIPTIVE FAULT TRAJECTORY', ['Fault / condition', 'Peak\nbacklog', 'Completed\nduring fault', 'Longest\ntransfer (s)', 'Normal in first\nwindow after'], [
    [labels[r['key']] + ' ' + ('R' if r['architecture'] == 'REST' else 'K'), r['maxSampledBacklog'], r['faultInterval']['completionsDuringFault'], f'{r["latency"]["max"]/1000:.1f}',
     'yes' if r['faultInterval']['firstWindowFullyAfterClearanceMeetsBaseline'] else 'no'] for r in data['runs'] if r['fault']], [.88, .50, .70, .62, .62])
para('Completed during fault counts transfers completed in the 30 s windows whose boundary fell inside the observed fault interval: one window for the adapter outage and two for the database delay. Longest transfer is the single slowest completion in the run. The last column states whether the first complete window entirely after fault clearance already met the predefined 90% goodput and 110% p95 baseline criteria. Under the predefined rule, every run scored about 59–60 s to backlog clearance and 119–120 s to recovery; no result was censored. Those scores are retained in the analysis package, but they reflect the window definition and its alignment, not a measured difference between implementations: in every run the systems had returned to baseline in the first eligible window and the rule\'s three-window confirmation set the recorded time.', size=8, indent=False)
heading('Resources and diagnostic context', True)
table('TABLE IV  TIME-WEIGHTED SERVICE RESOURCES', ['Condition', 'CPU\n(core eq.)', 'Memory\n(MiB)', 'App CPU\n(core eq.)'], [
    [labels[r['key']] + ' ' + ('R' if r['architecture'] == 'REST' else 'K'), f'{r["resources"]["means"]["cpuCores"]:.3f}', f'{r["resources"]["means"]["memoryMiB"]:.1f}', f'{r["resources"]["means"]["appCpuCores"]:.3f}'] for r in data['runs']], [1.14, .65, .75, .78])
para(f'In the no-fault pair, sampled service CPU was {sr["resources"]["means"]["cpuCores"]:.3f} core equivalents for REST and {sk["resources"]["means"]["cpuCores"]:.3f} for Kafka; memory was {sr["resources"]["means"]["memoryMiB"]:.1f} and {sk["resources"]["means"]["memoryMiB"]:.1f} MiB. The Kafka broker alone averaged {sk["resources"]["means"]["brokerMemoryMiB"]:.1f} MiB. Application-only CPU was also higher for Kafka in this pair. The same direction appears in both fault pairs, but resource use has no independent run-level uncertainty estimate. Sampling covered all but the first and last one to two seconds of each measurement.')
para('Diagnostic aggregates around the no-fault measurement show provider execution averaging about 30.1 ms for REST and 29.0 ms for Kafka. These aggregates include setup/boundary operations and are not exact measured-cohort estimates. Their similarity is consistent with a common provider setting, but it does not explain the remaining latency difference. Database calls, event dispatch, broker delivery, and scheduling overlap or nest; adding their quantiles would not produce a valid decomposition.')
para(f'Statement counts are additive and can be compared. In the no-fault screen the Kafka application issued {sk["statements"]["total"]:,} database statements against {sr["statements"]["total"]:,} for REST, including {sk["statements"]["poolStatements"]:,} pool-level statements, mostly inbox/outbox polling and dispatch bookkeeping, against {sr["statements"]["poolStatements"]:,}. The fault runs show the same ratio. An added 100 ms per database round trip therefore multiplies through roughly {sk["statements"]["total"]/sr["statements"]["total"]:.1f} times as many statements per run in Kafka, which is consistent with its larger in-fault stall under database delay, although it is an association, not a causal decomposition.')
heading('Supplementary archived no-fault pairs', True)
para(f'Before the reporting set, {words[len(same_pairs)]} REST–Kafka pairs were collected at 4 operations/s without faults under earlier experiment-harness revisions. Their frozen source archives differ from the reporting-set archive only in harness, protocol, and documentation files; the application-defining files are identical apart from a strict-versus-inclusive inequality in the calibration k6 check threshold, which does not affect traffic. All {words[len(same_app)//2]} pairs qualified with zero failures, zero pending transfers, and zero backlog at every 30 s boundary. Table V lists them with their realized order; they were not selected by outcome.')
table('TABLE V  ARCHIVED 4 OPS/S NO-FAULT PAIRS, SAME APPLICATION FILES', ['Pair', 'Order', 'Warm/\nmeas.', 'Arch.', 'n', 'Median\n(ms)', 'p95\n(ms)', 'p99\n(ms)'], [
    [r['label'].replace('Calibration pair', 'Calib.').replace('Screening pair', 'Screen.'), '1st' if r['firstInPair'] else '2nd', f"{r['warmupDuration'].rstrip('m')}/{r['measurementDuration'].rstrip('m')}", 'R' if r['architecture'] == 'REST' else 'K',
     f"{r['counts']['completed']:,}", ms_display(r['latency']['median']), ms_display(r['latency']['p95']), ms_display(r['latency']['p99'])] for r in same_app], [.56, .36, .44, .36, .42, .40, .40, .40])
para(f'Across these pairs REST p95 was {rng([r["latency"]["p95"] for r in same_r])} ms and Kafka p95 was {rng([r["latency"]["p95"] for r in same_k])} ms, with {words[len(kafka_first_pairs)]} of the three pairs executed Kafka first. The direction and size of the no-fault difference therefore repeated under both realized orders and at two run lengths, and the 8 operations/s screen fell in the same range. One further archived pair under an earlier application revision showed the same direction and is not tabulated. These are not a designed replication: harness revisions, dates, and protocols differ, and no confidence interval is derived from them.', size=8, indent=False)

heading('VI DISCUSSION AND LIMITATIONS')
para(f'The strongest supported statement is that REST had lower observed completion latency and sampled resource use in every pair, while both implementations preserved the tested ledger invariants and completed all admitted work. Neither demonstrated a capacity advantage. The archived pairs show the no-fault difference repeating under both orders and two run lengths; the fresh pairs do the same for the faults: under database delay, completions during the fault were {compare(infault(fdr), infault(fdk), "in-fault completions")} in the fresh pair, compared with {compare(infault(dr), infault(dk), "in-fault completions")} in the original pair. Two pairs per scenario are still not a replication estimate.')
para('The comparison bundles direct-call coordination with one durability path and event-driven coordination with another. Kafka adds inbox/outbox writes, broker communication, and dispatch scheduling. The 50 ms polling setting is a plausible contributor to its latency: a transfer crosses several event hops, and the observed median gap of about 90 ms is of the order of a few hops each waiting on average half of the polling interval, but these data cannot identify how much delay it caused. The statement-count ratio offers a separate, additive explanation for the database-delay result. A polling-interval sensitivity experiment would be needed before any causal claim. Co-located handlers and the shared ledger also remove distributed-service behaviours that may favor other designs.')
para('Internal validity is limited by one execution per architecture per scenario and order, a rebuilt implementation identity for the fresh pairs (application files identical apart from the package name), and a shared laptop/virtual-machine environment. Isolation checks remove competing containers but do not eliminate host scheduling, thermal, storage, or clock effects. Diagnostic collection adds overhead. The post-commit observation can be lost in a narrow crash interval; a ledger success without that observation must not be assigned a reconstructed latency.')
para('External validity is limited by deterministic arrivals, synthetic provider responses, one status-read target, limited account/amount patterns, and a single broker with replication factor one. The workload does not represent an empirically sampled operator trace. The adapter test does not cover application termination, broker loss, host failure, network partitions, or real provider-side economic reconciliation. No claim of general crash safety or production mobile-money performance follows.')
para('The publication scope was selected after the observations, and we preserve that history rather than relabel the data as a preregistered confirmatory study. The unequal no-fault and fault rates and durations prevent a controlled cross-scenario contrast. Further run-level repetition of every condition under both orders is the next priority; workload and polling sensitivities address different questions and should remain separate. The window-based recovery rule should also be revised before any confirmatory study, since it produced identical scores for visibly different trajectories.')

heading('VII ETHICS AND REPRODUCIBILITY')
para(f'All account identifiers, amounts, and provider responses are synthetic; no customer dataset or real funds are used. Local artifacts retain the evidence for all sixteen runs reported here (six reporting-set, four fresh, six archived), exact protocol/source/image identities, client attempts, transfer snapshots, traces, ledger checks, fault timings, and cleanup records. The analysis recomputes outcomes and verifies file inventories before and after reading. Tables and figures are generated from the resulting JSON rather than manually transcribed. The reporting-set source identity is 7995e552…, the fresh-pair identity is {data["freshFaultSourceSha256"][:8]}…; full hashes, pinned image identities, and review receipts are in the results package.')
para(f'The code, a redacted copy of the evidence for all sixteen runs, the analysis, the figures, and this draft are publicly available [17] under Apache-2.0 for code and CC BY 4.0 for evidence and text, {artifact_cite}. The public copy is generated from the private archive by scripts that replace the capture host name, home paths, Docker identity, and the names of unrelated containers with neutral tokens and record the SHA-256 of every original file, so it can be verified against the archive; no experiment number depends on a redacted field. The public repository also lets another site start its own guarded chain with the same protocols. Independent reproduction by a third party has not yet been established.')

heading('VIII CONCLUSION')
para(f'This exploratory benchmark links synthetic transfer traffic to durable completion and ledger correctness under one common asynchronous client contract. In sixteen executions (six reporting-set, six archived same-application, four fresh Kafka-first), every measured transfer completed correctly. REST had lower observed completion tails and sampled resource use in every pair, under both realized orders; both implementations accumulated and later cleared fault-induced backlogs. While database responses were delayed, in-fault completions were {compare(infault(dr), infault(dk), "in-fault completions")} in the original pair and {compare(infault(fdr), infault(fdk), "in-fault completions")} in the fresh Kafka-first pair, consistent in both pairs with Kafka\'s larger statement count. The results support a bounded comparison of these implementations and identify questions for further testing. They do not establish capacity, general architecture superiority, or equivalent recovery time. Further repetitions under both orders and focused configuration sensitivity would strengthen those inferences.')

heading('ACKNOWLEDGMENTS')
para('The author used Claude to improve the readability and language quality of this manuscript, and Claude Code for code development support. All content was reviewed and edited by the author, who takes full responsibility for the final work.', indent=False)

heading('REFERENCES')
references = [
    '[1] GSMA, “P2P Transfers,” Mobile Money API Developer Portal. Accessed Sep. 21, 2026. https://developer.mobilemoneyapi.io/use-cases/p-2-p-transfers/',
    '[2] J. Kazanavičius and D. Mažeika, “The Evaluation of Microservice Communication While Decomposing Monoliths,” Computing and Informatics, vol. 42, no. 1, pp. 1–36, 2023. doi: 10.31577/cai_2023_1_1.',
    '[3] S.-B. Son, C. Lee, and K. Lee, “Comparison of Distributed Transactions in Microservice Architecture,” The Journal of Information Technology and Architecture, vol. 20, no. 4, pp. 281–294, 2023. doi: 10.22865/jita.2023.20.4.281.',
    '[4] H. Kristianto and A. Zahra, “Performance Analysis of Choreography and Orchestration in Microservices Architecture,” Journal of Theoretical and Applied Information Technology, vol. 99, no. 18, pp. 4220–4230, 2021. https://www.jatit.org/volumes/Vol99No18/4Vol99No18.pdf',
    '[5] W. Hasselbring, “Benchmarking as Empirical Standard in Software Engineering Research,” in Proc. Evaluation and Assessment in Software Engineering (EASE), pp. 365–372, 2021. doi: 10.1145/3463274.3463361.',
    '[6] T. Kalibera and R. Jones, “Rigorous Benchmarking in Reasonable Time,” in Proc. International Symposium on Memory Management (ISMM), pp. 63–74, 2013. doi: 10.1145/2464157.2464160.',
    '[7] Grafana Labs, “Constant arrival rate,” k6 documentation. Accessed Sep. 21, 2026. https://grafana.com/docs/k6/latest/using-k6/scenarios/executors/constant-arrival-rate/',
    '[8] Docker, “docker container stats,” Docker documentation. Accessed Sep. 21, 2026. https://docs.docker.com/reference/cli/docker/container/stats/',
    '[9] A. Nadeem and M. Z. Malik, “A Case for Microservices Orchestration Using Workflow Engines,” in Proc. IEEE/ACM 44th Int. Conf. Software Engineering: New Ideas and Emerging Results (ICSE-NIER), pp. 6–10, 2022. doi: 10.1145/3510455.3512777.',
    '[10] P. Hegde N et al., “An Empirical Evaluation of Synchronous vs Asynchronous Microservice Architectures for Enterprise Messaging Platforms,” in Proc. Int. Conf. Artificial Intelligence and Data Engineering (AIDE), pp. 663–666, 2026. doi: 10.1109/AIDE69088.2026.11545016.',
    '[11] S. Aydın and C. B. Çebi, “Comparison of Choreography vs Orchestration Based Saga Patterns in Microservices,” in Proc. Int. Conf. Electrical, Computer and Energy Technologies (ICECET), pp. 1–6, 2022. doi: 10.1109/ICECET55527.2022.9872665.',
    '[12] C. Richardson, Microservices Patterns: With Examples in Java. Shelter Island, NY, USA: Manning, 2018.',
    '[13] M. Bianchi, M. Bouvard, R. Gomes, A. Rhodes, and V. Shreeti, “Mobile payments and interoperability: Insights from the academic literature,” Information Economics and Policy, vol. 65, art. 101068, 2023. doi: 10.1016/j.infoecopol.2023.101068.',
    '[14] Committee on Payments and Market Infrastructures, “Promoting the harmonisation of application programming interfaces to enhance cross-border payments: Recommendations and toolkit,” Bank for International Settlements, Basel, 2024. https://www.bis.org/cpmi/publ/d224.htm',
    '[15] GSMA, “Mobile Money API Specification 1.2.0: Fundamentals,” GSMA, London, 2021. https://www.gsma.com/mobilefordevelopment/wp-content/uploads/2021/10/Mobile-Money-API-Specification-1.2.0-Fundamentals.pdf',
    '[16] Mojaloop Foundation, “Mojaloop Hub,” Mojaloop documentation, 2022. https://docs.mojaloop.io/technical/overview/',
    f'[17] S. Lubanga, “Mobile-money coordination benchmark: Code, redacted evidence, analysis and paper,” {ARTIFACT_VERSION},{artifact_ref} {ARTIFACT_URL}'
]
for ref in references:
    p = para(ref, size=8, indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(.16); p.paragraph_format.first_line_indent = Inches(-.16)
    p.paragraph_format.space_after = Pt(2)

doc.core_properties.title = title
doc.core_properties.subject = 'Results draft from six reporting-set runs, six archived same-application no-fault runs and four Kafka-first fault runs'
doc.core_properties.author = 'Shaban Lubanga'
doc.core_properties.comments = 'Revised copy. One execution per architecture per condition; not a confirmatory study.'
doc.save(target)
assert hashlib.sha256(original.read_bytes()).hexdigest() == source_hash
(ART / 'manuscript.md').write_text('\n'.join(manuscript))
(ART / 'document-build.json').write_text(json.dumps({
    'original': original.name, 'originalSha256': source_hash,
    'output': str(target.relative_to(ROOT)), 'outputSha256': hashlib.sha256(target.read_bytes()).hexdigest(),
    'analysisSha256': hashlib.sha256((ART / 'analysis.json').read_bytes()).hexdigest(),
    'originalUnchanged': True, 'pythonDocxVersion': '1.2.0',
    'scientificStatus': 'Exploratory draft; publication readiness and replication not established'
}, indent=2) + '\n')
print(target)
